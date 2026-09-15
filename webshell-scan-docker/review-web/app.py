#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Local review portal for webshell scan findings.

The app is intentionally dependency-light: stdlib HTTP + SQLite, with optional
openpyxl/yara support when they are installed in the scanner environment.
"""

from __future__ import annotations

import argparse
import base64
import binascii
import csv
import hashlib
import hmac
import html
import io
import json
import mimetypes
import os
import posixpath
import re
import secrets
import shutil
import sqlite3
import sys
import textwrap
import time
import urllib.parse
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from http.cookies import SimpleCookie
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path, PurePosixPath

try:
    import openpyxl
except Exception:  # pragma: no cover - optional runtime dependency
    openpyxl = None

try:
    import yara
except Exception:  # pragma: no cover - optional runtime dependency
    yara = None


APP_DIR = Path(__file__).resolve().parent
SCANNER_DIR = APP_DIR.parent
DATA_DIR = APP_DIR / "data"
DB_PATH = DATA_DIR / "review.db"
CONFIG_PATH = DATA_DIR / "config.json"
MAX_UPLOAD_BYTES = 512 * 1024 * 1024
MAX_VIEW_BYTES = 2 * 1024 * 1024
YARA_TIMEOUT_SECONDS = 20
FINDING_FILE_EXTENSIONS = (".csv", ".xlsx", ".xlsm")
SESSION_COOKIE = "review_session"
ADMIN_COOKIE = "review_admin"
SESSION_TTL_SECONDS = 8 * 60 * 60
VALID_STATUSES = ("pending", "reviewing", "confirmed", "benign", "remediated")

sys.path.insert(0, str(SCANNER_DIR))

try:
    from libs.deobfuscate import iter_deobfuscated_layers, should_deobfuscate
except Exception:  # pragma: no cover - scanner tree may be incomplete
    iter_deobfuscated_layers = None
    should_deobfuscate = None

try:
    from libs.review_bundle import (
        ReviewBundleError,
        bundle_passphrase,
        decrypt_review_bundle_bytes,
    )
except Exception:  # pragma: no cover - scanner tree may be incomplete
    ReviewBundleError = RuntimeError
    bundle_passphrase = None
    decrypt_review_bundle_bytes = None


REVIEW_JS = b"""(function(){
  document.addEventListener('DOMContentLoaded', function(){
    var toggle = document.querySelector('[data-rail-toggle]');
    var rail = document.querySelector('[data-rail]');
    if (toggle && rail) {
      var saved = localStorage.getItem('rail_expanded');
      if (saved === 'true') { rail.classList.add('expanded'); toggle.setAttribute('aria-expanded', 'true'); }
      toggle.addEventListener('click', function(){
        var exp = rail.classList.toggle('expanded');
        toggle.setAttribute('aria-expanded', exp ? 'true' : 'false');
        localStorage.setItem('rail_expanded', exp ? 'true' : 'false');
      });
    }
    document.addEventListener('click', function(e){
      var btn = e.target.closest('[data-copy-value]');
      if (!btn) return;
      var val = btn.getAttribute('data-copy-value');
      if (val && navigator.clipboard) {
        navigator.clipboard.writeText(val).then(function(){
          var orig = btn.textContent;
          btn.textContent = 'Copied!';
          setTimeout(function(){ btn.textContent = orig; }, 2000);
        }).catch(function(){});
      }
    });
  });
})();"""


@dataclass
class UploadedPart:
    name: str
    filename: str | None
    content_type: str
    data: bytes


def now_ts() -> int:
    return int(time.time())


def read_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def ensure_dirs() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    (DATA_DIR / "cases").mkdir(parents=True, exist_ok=True)
    if not CONFIG_PATH.exists():
        CONFIG_PATH.write_text(
            json.dumps({"admin_key": secrets.token_urlsafe(32)}, indent=2),
            encoding="utf-8",
        )


def get_admin_key() -> str:
    env_key = os.environ.get("REVIEW_ADMIN_KEY", "").strip()
    if env_key:
        return env_key
    return str(read_json(CONFIG_PATH, {}).get("admin_key", ""))



USER_COOKIE = "review_user"


def login_page(message: str = "") -> bytes:
    notice = f'<div class="warn">{html.escape(message)}</div>' if message else ""
    body = f"""
    <section class="login-shell analyst-shell">
      <div class="login-brand">
        <div class="eyebrow">STAGING · Authorized access only</div>
        <h1>Webshell Review</h1>
        <p>Evidence-led triage for webshell findings, gray-zone candidates, and remediation status.</p>
        <div class="muted">Only provisioned accounts may enter this workspace.</div>
      </div>
      <div class="login-card">
        <h2>Đăng nhập tài khoản</h2>
        {notice}
        <form method="post" action="/login" class="stack">
          <label>Tên đăng nhập<br><input name="username" autocomplete="username" required style="width:100%"></label>
          <label>Mật khẩu<br><input name="password" type="password" autocomplete="current-password" required style="width:100%"></label>
          <button type="submit">Đăng nhập</button>
        </form>
      </div>
    </section>
    """
    return html_page("Đăng nhập", body)

def normalize_username(value: str) -> str:
    return " ".join(str(value or "").strip().lower().split())


def hash_password(password: str, minimum_length: int = 12) -> str:
    if not isinstance(password, str) or len(password) < int(minimum_length):
        raise ValueError("Password is shorter than the configured minimum")
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(
        password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32,
    )
    encode = lambda value: base64.urlsafe_b64encode(value).decode("ascii").rstrip("=")
    return "scrypt$v1$%s$%s" % (encode(salt), encode(digest))


def verify_password(password: str, encoded: str) -> bool:
    try:
        version, marker, salt_text, digest_text = str(encoded or "").split("$", 3)
        if version != "scrypt" or marker != "v1":
            return False
        pad = lambda value: value + "=" * (-len(value) % 4)
        salt = base64.urlsafe_b64decode(pad(salt_text))
        expected = base64.urlsafe_b64decode(pad(digest_text))
        actual = hashlib.scrypt(
            str(password).encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=len(expected),
        )
        return hmac.compare_digest(actual, expected)
    except (TypeError, ValueError, binascii.Error):
        return False


@contextmanager
def connect_db() -> Iterator[sqlite3.Connection]:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def ensure_audit_log_schema(conn: sqlite3.Connection) -> None:
    columns = {row["name"]: row for row in conn.execute("PRAGMA table_info(audit_log)")}
    case_column = columns.get("case_id")
    if not case_column or not case_column["notnull"]:
        return
    conn.execute("ALTER TABLE audit_log RENAME TO audit_log_legacy")
    conn.execute(
        """
        CREATE TABLE audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            case_id INTEGER,
            finding_id INTEGER,
            action TEXT NOT NULL,
            detail TEXT,
            created_at INTEGER NOT NULL
        )
        """
    )
    conn.execute(
        "INSERT INTO audit_log(id,case_id,finding_id,action,detail,created_at) "
        "SELECT id,case_id,finding_id,action,detail,created_at FROM audit_log_legacy"
    )
    conn.execute("DROP TABLE audit_log_legacy")

def init_db() -> None:
    with connect_db() as conn:
        conn.executescript(
            """
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS cases (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                share_key TEXT NOT NULL UNIQUE,
                viewer_key TEXT UNIQUE,
                created_at INTEGER NOT NULL,
                source_zip_name TEXT,
                finding_count INTEGER NOT NULL DEFAULT 0,
                source_count INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS findings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id INTEGER NOT NULL,
                file_path TEXT NOT NULL,
                sha256 TEXT,
                score INTEGER,
                scan_time TEXT,
                description TEXT,
                extension TEXT,
                candidate_kind TEXT NOT NULL DEFAULT 'finding',
                status TEXT NOT NULL DEFAULT 'pending',
                reviewer_note TEXT NOT NULL DEFAULT '',
                local_relpath TEXT,
                has_source INTEGER NOT NULL DEFAULT 0,
                created_at INTEGER NOT NULL,
                FOREIGN KEY(case_id) REFERENCES cases(id)
            );
            CREATE INDEX IF NOT EXISTS idx_findings_case ON findings(case_id);
            CREATE INDEX IF NOT EXISTS idx_findings_status ON findings(case_id, status);
            CREATE TABLE IF NOT EXISTS layers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                finding_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                relpath TEXT NOT NULL,
                size INTEGER NOT NULL,
                sha256 TEXT NOT NULL,
                UNIQUE(finding_id, name),
                FOREIGN KEY(finding_id) REFERENCES findings(id)
            );
            CREATE TABLE IF NOT EXISTS evidence (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                finding_id INTEGER NOT NULL,
                layer_name TEXT NOT NULL,
                rule_name TEXT NOT NULL,
                string_id TEXT,
                start_offset INTEGER,
                end_offset INTEGER,
                matched_text TEXT,
                score INTEGER,
                description TEXT,
                reference TEXT,
                author TEXT,
                FOREIGN KEY(finding_id) REFERENCES findings(id)
            );
            CREATE INDEX IF NOT EXISTS idx_evidence_finding ON evidence(finding_id);
            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id INTEGER,
                finding_id INTEGER,
                action TEXT NOT NULL,
                detail TEXT,
                created_at INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                case_id INTEGER NOT NULL,
                role TEXT NOT NULL DEFAULT 'reviewer',
                created_at INTEGER NOT NULL,
                expires_at INTEGER NOT NULL,
                FOREIGN KEY(case_id) REFERENCES cases(id)
            );
            CREATE INDEX IF NOT EXISTS idx_sessions_case ON sessions(case_id);
            CREATE TABLE IF NOT EXISTS admin_sessions (
                token TEXT PRIMARY KEY,
                csrf_token TEXT NOT NULL,
                user_id INTEGER,
                created_at INTEGER NOT NULL,
                expires_at INTEGER NOT NULL,
                FOREIGN KEY(user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'user' CHECK(role IN ('admin', 'user')),
                active INTEGER NOT NULL DEFAULT 1,
                must_change_password INTEGER NOT NULL DEFAULT 0,
                created_at INTEGER NOT NULL,
                updated_at INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS user_sessions (
                token TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL,
                csrf_token TEXT NOT NULL,
                created_at INTEGER NOT NULL,
                expires_at INTEGER NOT NULL,
                FOREIGN KEY(user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS case_assignments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,
                assigned_by_user_id INTEGER,
                created_at INTEGER NOT NULL,
                revoked_at INTEGER,
                FOREIGN KEY(case_id) REFERENCES cases(id),
                FOREIGN KEY(user_id) REFERENCES users(id)
            );
            CREATE INDEX IF NOT EXISTS idx_case_assignments_active ON case_assignments(case_id, user_id, active);
            CREATE TABLE IF NOT EXISTS upload_credentials (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                key_hash TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,
                created_at INTEGER NOT NULL,
                rotated_at INTEGER,
                created_by_user_id INTEGER,
                FOREIGN KEY(created_by_user_id) REFERENCES users(id)
            );
            """
        )
        ensure_admin_session_schema(conn)
        ensure_column(conn, "admin_sessions", "user_id", "INTEGER")
        ensure_column(conn, "cases", "viewer_key", "TEXT")
        ensure_column(conn, "cases", "uploaded_by_user_id", "INTEGER")
        ensure_column(conn, "cases", "upload_origin", "TEXT NOT NULL DEFAULT 'admin_web'")
        ensure_column(conn, "sessions", "role", "TEXT NOT NULL DEFAULT 'reviewer'")
        ensure_column(conn, "findings", "candidate_kind", "TEXT NOT NULL DEFAULT 'finding'")
        ensure_column(conn, "users", "must_change_password", "INTEGER NOT NULL DEFAULT 0")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_findings_kind ON findings(case_id, candidate_kind)")
        conn.execute(
            "UPDATE findings SET candidate_kind='finding' "
            "WHERE candidate_kind IS NULL OR candidate_kind NOT IN ('finding', 'gray')"
        )
        conn.execute(
            "UPDATE cases SET upload_origin='admin_web' "
            "WHERE upload_origin IS NULL OR upload_origin NOT IN ('admin_web', 'remote_upload')"
        )
        for row in conn.execute("SELECT id FROM cases WHERE viewer_key IS NULL OR viewer_key=''", ()):
            conn.execute("UPDATE cases SET viewer_key=? WHERE id=?", (secrets.token_urlsafe(24), row["id"]))
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_cases_viewer_key ON cases(viewer_key)")
        if conn.execute("SELECT 1 FROM users LIMIT 1").fetchone() is None:
            bootstrap_password = os.environ.get("REVIEW_BOOTSTRAP_PASSWORD", "")
            if bootstrap_password:
                created = now_ts()
                conn.execute(
                    "INSERT INTO users(username,password_hash,role,active,must_change_password,created_at,updated_at) "
                    "VALUES(?,?,?,?,?,?,?)",
                    ("antt", hash_password(bootstrap_password, minimum_length=8), "admin", 1, 1, created, created),
                )
            else:
                # No default password is created without explicit out-of-band bootstrap secret.
                pass

        ensure_audit_log_schema(conn)


def ensure_column(conn: sqlite3.Connection, table: str, column: str, definition: str) -> None:
    existing = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}
    if column not in existing:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")



def ensure_admin_session_schema(conn: sqlite3.Connection) -> None:
    ensure_column(conn, "admin_sessions", "csrf_token", "TEXT NOT NULL DEFAULT ''")


def scanner_policy() -> dict:
    return read_json(SCANNER_DIR / "config" / "scan-policy.json", {})


def html_page(title: str, body: str, *, refresh: bool = False) -> bytes:
    refresh_tag = '<meta http-equiv="refresh" content="5">' if refresh else ""
    admin_shell = body.startswith("<!-- soc-admin -->")
    body = body.removeprefix("<!-- soc-admin -->")
    workspace_links = (
        '<a class="soc-nav-item active" href="/" title="Dashboard tổng quan"><span class="nav-icon">▦</span><span class="nav-text">Dashboard</span></a>'
        '<a class="soc-nav-item" href="/admin" title="All reports"><span class="nav-icon">▣</span><span class="nav-text">Reports</span></a>'
        '<a class="soc-nav-item" href="/user" title="Assigned to me"><span class="nav-icon">◈</span><span class="nav-text">Assigned</span></a>'
        '<a class="soc-nav-item" href="/?kind=gray" title="Gray zone"><span class="nav-icon">◇</span><span class="nav-text">Gray zone</span></a>'
        if admin_shell
        else '<a class="soc-nav-item active" href="/" title="Dashboard tổng quan"><span class="nav-icon">▦</span><span class="nav-text">Dashboard</span></a>'
        '<a class="soc-nav-item" href="/user" title="Assigned to me"><span class="nav-icon">◈</span><span class="nav-text">Assigned</span></a>'
    )
    admin_links = (
        '<div class="soc-sidebar-label">ADMIN</div>'
        '<a class="soc-nav-item" href="/admin#users" title="Users"><span class="nav-icon">♙</span><span class="nav-text">Users</span></a>'
        '<a class="soc-nav-item" href="/admin#import" title="Import"><span class="nav-icon">📥</span><span class="nav-text">Import</span></a>'
        if admin_shell
        else ""
    )
    sidebar = f'<aside class="soc-sidebar" data-rail><button class="rail-toggle-btn" data-rail-toggle aria-expanded="false" title="Mở/đóng rail">☰</button><div class="soc-sidebar-label">WORKSPACE</div>{workspace_links}{admin_links}</aside>'
    page = f"""<!doctype html>
<html lang="vi">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">{refresh_tag}<title>{html.escape(title)}</title>
<style>
:root {{ color-scheme:light; --canvas:#F5F7FB; --surface:#fff; --surface-muted:#F7F9FC; --primary:#2457D6; --primary-soft:#E8F0FF; --navy:#183B70; --text:#24324A; --text-strong:#18243A; --muted:#718096; --muted-readable:#52657C; --line:#E5EAF2; --line-strong:#C9D3E2; --danger:#D64545; --danger-soft:#FDECEC; --warning:#B7791F; --warning-soft:#FFF6D8; --success:#25855A; --success-soft:#E7F6EE; --code:#111827; --code-text:#E6EDF7; --code-muted:#AFC0D6; --code-line:#24324A; --focus:#2563EB; --mark:#fff2a8; --mark2:#b8f2e6; --mark3:#ffd6a5; }}
*{{box-sizing:border-box}} body{{margin:0;background:var(--canvas);color:var(--text-strong);font:15px/1.6 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}} a{{color:var(--primary);text-decoration:none}}a:hover{{text-decoration:underline}} .muted,.eyebrow{{color:var(--muted-readable)}} .detail-meta .path{{color:var(--text-strong);font-size:16px;line-height:1.45;overflow-wrap:anywhere;word-break:break-word}} .detail-meta{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:20px 22px;margin-bottom:16px}} .detail-meta .filters{{gap:8px;margin:12px 0}} .detail-meta .meta-actions{{display:flex;align-items:center;gap:10px;flex-wrap:wrap}} .detail-meta .meta-actions .path{{flex:1 1 420px;min-width:0;font-size:12px}} .detail-sidebar{{position:sticky;top:18px;max-height:calc(100vh - 36px);overflow:auto}} .detail-sidebar h2{{color:var(--text-strong);font-size:16px;margin-top:4px}} .detail-sidebar label{{font-size:14px}} .detail-sidebar textarea{{min-height:110px;resize:vertical;width:100%}} .detail-viewer{{min-width:0;padding:14px}} .finding-actions,.action-group{{display:flex;gap:8px;flex-wrap:wrap;margin:12px 0 4px}} .finding-actions button,.action-group button{{min-height:40px}} .action-confirm{{background:var(--danger);border-color:var(--danger)}} .action-remediate{{background:var(--success);border-color:var(--success)}} .action-benign{{background:#fff;color:var(--muted);border-color:var(--line-strong)}}
.soc-header{{background:#fff;color:var(--text-strong);min-height:56px;border-bottom:1px solid var(--line);box-shadow:0 1px 3px rgba(0,0,0,.04)}}.soc-header-inner{{max-width:1440px;margin:auto;padding:0 20px;min-height:56px;display:flex;justify-content:space-between;align-items:center;gap:16px}}.soc-brand{{display:flex;align-items:center;gap:10px;color:var(--text-strong)}}.soc-brand-mark{{width:30px;height:30px;border:2px solid var(--primary);color:var(--primary);border-radius:50%;display:grid;place-items:center;font-size:11px;font-weight:800}}.soc-brand-title{{font-size:15px;font-weight:800;line-height:1.1;display:block}}.soc-brand-subtitle{{font-size:11px;color:var(--muted);display:block}}.soc-context{{display:flex;align-items:center;gap:10px}}.soc-env{{background:var(--primary-soft);color:var(--primary);border:1px solid #b8d2ff;border-radius:6px;padding:4px 8px;font-size:11px;font-weight:800}}.soc-user{{width:28px;height:28px;border-radius:50%;background:var(--primary);color:#fff;display:grid;place-items:center;font-size:12px;font-weight:800}}
.soc-layout{{display:grid;grid-template-columns:52px minmax(0,1fr);min-height:calc(100vh - 56px);max-width:1440px;margin:auto;transition:grid-template-columns .2s ease}}.soc-sidebar{{background:#fff;border-right:1px solid var(--line);padding:12px 6px;overflow:hidden}}.soc-sidebar.expanded{{grid-template-columns:220px minmax(0,1fr)}}.soc-layout:has(.soc-sidebar.expanded){{grid-template-columns:220px minmax(0,1fr)}}.rail-toggle-btn{{background:transparent;border:0;color:var(--muted);width:100%;padding:8px 0;text-align:center;cursor:pointer;font-size:16px;border-radius:6px;margin-bottom:12px}}.rail-toggle-btn:hover{{background:var(--surface-muted);color:var(--text-strong)}}.soc-sidebar-label{{padding:6px 6px 4px;color:#8a99ad;font-size:9px;font-weight:800;letter-spacing:.12em;white-space:nowrap;overflow:hidden}}.soc-nav-item{{display:flex;align-items:center;gap:10px;padding:9px 10px;margin-bottom:3px;border-radius:7px;color:var(--muted-readable);font-size:13px;white-space:nowrap;overflow:hidden}}.soc-nav-item .nav-icon{{font-size:15px;min-width:20px;text-align:center}}.soc-nav-item:hover,.soc-nav-item.active{{background:var(--primary-soft);color:var(--primary);text-decoration:none;font-weight:750}}.soc-content{{min-width:0;padding:22px 26px}}
.band,.card{{background:#fff;border:1px solid var(--line);border-radius:9px;box-shadow:0 2px 8px rgba(31,55,86,.04)}}.band{{padding:20px;margin-bottom:16px}}h1,h2{{margin:0 0 12px;line-height:1.2}}h1{{font-size:24px;color:var(--text)}}h2{{font-size:17px;color:var(--navy)}}.eyebrow{{color:var(--muted);font-size:11px;font-weight:800;letter-spacing:.12em;text-transform:uppercase;margin-bottom:8px}}.muted{{color:var(--muted)}}
.hero{{background:#fff;color:var(--text-strong);border:1px solid var(--line);box-shadow:0 2px 6px rgba(0,0,0,.03)}}.hero h1{{color:var(--text-strong)}}.hero .muted,.hero .eyebrow{{color:var(--muted-readable)}}.grid,.kpi-row{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:13px}}.metric,.kpi{{background:var(--surface-muted);border:1px solid var(--line);border-radius:8px;padding:13px;min-height:82px;position:relative}}.metric::before,.kpi::before{{content:"";position:absolute;inset:0 auto auto 0;width:100%;height:3px;background:var(--primary)}}.metric-main::before,.kpi.danger::before{{background:var(--danger)}}.metric-gray::before,.kpi.warn::before{{background:var(--warning)}}.metric-ok::before,.kpi.ok::before{{background:var(--success)}}.metric b,.kpi b{{display:block;color:var(--navy);font-size:27px;line-height:1.1;margin-bottom:5px}}.metric span,.kpi span{{color:var(--muted);font-size:12px}}
input,select,textarea,button{{font:inherit;border:1px solid var(--line-strong);border-radius:6px;background:#fff;padding:9px 11px;color:var(--text)}}label{{color:var(--text);font-weight:700}}button,.button{{background:var(--primary);border-color:var(--primary);color:#fff;cursor:pointer;font-weight:800}}button:hover,.button:hover{{background:#1D4ED8;text-decoration:none}}button.secondary,.button.secondary{{background:#fff;color:var(--primary);border-color:var(--line-strong)}}input:focus-visible,select:focus-visible,textarea:focus-visible,button:focus-visible,a:focus-visible{{outline:3px solid #9cc2ff;outline-offset:2px}}
.filters{{display:flex;flex-wrap:wrap;gap:7px;align-items:center}}.tablewrap{{overflow-x:auto}}table{{width:100%;border-collapse:separate;border-spacing:0;min-width:720px}}th,td{{padding:11px 10px;border-bottom:1px solid var(--line);vertical-align:top}}th{{text-align:left;background:var(--surface-muted);color:var(--muted);font-size:11px;text-transform:uppercase;letter-spacing:.06em}}tbody tr:hover td{{background:#F3F7FF}}.path{{font:12px/1.45 ui-monospace,SFMono-Regular,Consolas,monospace;overflow-wrap:anywhere}}.pill,.kind-finding,.kind-gray,.source-ok,.source-missing,.status-confirmed,.status-remediated,.status-benign,.status-pending,.status-reviewing{{display:inline-flex;align-items:center;border-radius:999px;padding:4px 9px;font-size:11px;font-weight:800}}.kind-finding{{color:var(--danger);background:var(--danger-soft)}}.kind-gray,.source-missing{{color:var(--warning);background:var(--warning-soft)}}.source-ok{{color:var(--success);background:var(--success-soft)}}.status-confirmed{{color:var(--danger);background:var(--danger-soft)}}.status-remediated{{color:var(--success);background:var(--success-soft)}}.status-benign{{color:var(--muted);background:#EDF2F7}}.status-pending,.status-reviewing{{color:var(--navy);background:var(--primary-soft)}}.pill{{color:var(--navy);background:#fff;border:1px solid var(--line-strong)}}.assignment-panel{{margin-top:6px;padding:7px 0}}.assignment-panel summary{{color:var(--primary);cursor:pointer;font-weight:800;font-size:12px}}.admin-header{{display:flex;align-items:center;justify-content:space-between;gap:24px;padding:8px 0 18px}}.admin-identity{{display:flex;align-items:center;gap:9px}}.action-bar{{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:16px}}.admin-header+ .action-bar{{margin-top:0}}.eyebrow{{color:var(--muted-readable)}}.muted{{color:var(--muted-readable)}}
.login-shell{{max-width:1050px;margin:58px auto;display:grid;grid-template-columns:1fr 420px;gap:28px;align-items:stretch}}.login-brand{{background:var(--navy);color:#fff;border-radius:10px;padding:40px;display:flex;flex-direction:column;justify-content:center;min-height:420px}}.login-brand h1{{color:#fff;font-size:36px;margin-top:10px}}.login-brand p{{color:#d7e6fb;max-width:430px}}.login-card{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:32px;box-shadow:0 8px 24px rgba(16,42,67,.08);align-self:center}}.login-card h2{{font-size:22px;margin-bottom:20px}}.login-card button{{width:100%;margin-top:6px}}
.split{{display:grid;grid-template-columns:350px minmax(0,1fr);gap:16px;align-items:start}}.codewrap{{background:var(--code);border-radius:8px;overflow:auto;max-height:74vh}}.code-toolbar{{background:#0b1220;color:var(--code-text);padding:9px 12px;border-bottom:1px solid #34495e;font:13px/1.45 ui-monospace,SFMono-Regular,Consolas,monospace}}.code{{margin:0;color:var(--code-text);font:13px/1.65 ui-monospace,SFMono-Regular,Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere;word-break:break-word}}.ln{{display:inline-block;width:58px;padding-right:12px;text-align:right;color:var(--code-muted);user-select:none}}mark{{color:#111827;padding:0 1px;border-radius:2px}}mark.m0{{background:var(--mark)}}mark.m1{{background:var(--mark2)}}mark.m2{{background:var(--mark3)}}.evidence-viewer{{min-width:0}}.detail-section{{border-top:1px solid var(--line);padding-top:12px;margin-top:14px}}.detail-section h2{{margin-bottom:8px}}.rule-list{{color:var(--text-strong);font-size:14px}}.rule-hit code{{overflow-wrap:anywhere;word-break:break-word}}.details summary{{cursor:pointer;color:var(--primary);font-weight:700}}.details .detail-text{{color:var(--muted-readable);overflow-wrap:anywhere}}
@media(max-width:980px){{.soc-layout{{grid-template-columns:1fr}}.soc-sidebar{{border-right:0;border-bottom:1px solid var(--line);display:flex;gap:4px;overflow:auto;padding:10px}}.soc-sidebar-label{{display:none}}.soc-nav-item{{white-space:nowrap}}.soc-content{{padding:18px 14px}}.grid,.kpi-row{{grid-template-columns:repeat(2,minmax(0,1fr))}}.split{{grid-template-columns:1fr}}.detail-sidebar{{position:static;max-height:none;overflow:visible}}.detail-meta .meta-actions .path{{flex-basis:100%}}.admin-header{{align-items:flex-start;flex-direction:column}}.login-shell{{grid-template-columns:1fr;margin:24px auto}}.login-brand{{min-height:260px}}}}@media(max-width:620px){{body{{font-size:15px}}.soc-header-inner{{padding:0 14px}}.soc-brand-subtitle{{display:none}}.grid,.kpi-row{{grid-template-columns:1fr}}.soc-content{{padding:14px 10px}}.login-shell{{margin:12px auto;gap:12px}}.login-brand,.login-card{{padding:24px}}.login-brand h1{{font-size:29px}}.detail-meta{{padding:16px}}.detail-meta .path{{font-size:14px}}.detail-viewer{{padding:8px}}.code{{font-size:12px;line-height:1.6}}.finding-actions button,.action-group button{{flex:1 1 145px}}table{{min-width:0}}thead{{display:none}}table,tbody,tr,td{{display:block;width:100%}}tbody tr{{margin-bottom:10px;border:1px solid var(--line);border-radius:8px;overflow:hidden}}td{{display:grid;grid-template-columns:100px minmax(0,1fr);gap:8px}}td::before{{content:attr(data-label);color:var(--muted);font-size:10px;font-weight:800;text-transform:uppercase}}}}
</style>
</head><body>
<header class="soc-header"><div class="soc-header-inner"><a class="soc-brand" href="/"><span class="soc-brand-mark">WR</span><span><b class="soc-brand-title">Webshell Review</b><small class="soc-brand-subtitle">Security Operations Center · Review Workspace</small></span></a><div class="soc-context"><span class="soc-env">STAGING</span><span class="soc-user">A</span></div></div></header>
<div class="soc-layout">{sidebar}<main class="soc-content">{body}</main></div>
<script src="/static/review.js"></script>
</body></html>"""
    return page.encode("utf-8")


def safe_int(value, default=0):
    try:
        if value in (None, ""):
            return default
        return int(float(str(value).strip()))
    except Exception:
        return default


def normalize_header(value: str) -> str:
    return re.sub(r"\s+", "_", str(value or "").strip().upper())


def parse_findings_file(filename: str, data: bytes) -> list[dict]:
    suffix = Path(filename or "").suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        if openpyxl is None:
            raise RuntimeError("Thiếu openpyxl nên chưa đọc được file Excel.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [normalize_header(v) for v in next(rows, [])]
        result = []
        for row in rows:
            item = {}
            for idx, header in enumerate(headers):
                if not header or header.startswith("UNNAMED"):
                    continue
                item[header] = row[idx] if idx < len(row) else None
            if item.get("FILE"):
                result.append(item)
        return result

    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    result = []
    for row in reader:
        item = {normalize_header(k): v for k, v in row.items() if k}
        if item.get("FILE"):
            result.append(item)
    return result


def parse_multipart(headers, body: bytes) -> dict[str, list[UploadedPart]]:
    content_type = headers.get("Content-Type", "")
    raw = b"Content-Type: " + content_type.encode("latin-1") + b"\r\nMIME-Version: 1.0\r\n\r\n" + body
    msg = BytesParser(policy=policy.default).parsebytes(raw)
    parts: dict[str, list[UploadedPart]] = {}
    if not msg.is_multipart():
        return parts
    for part in msg.iter_parts():
        cd = part.get("Content-Disposition", "")
        if "form-data" not in cd:
            continue
        name = part.get_param("name", header="content-disposition")
        if not name:
            continue
        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""
        parts.setdefault(name, []).append(
            UploadedPart(name=name, filename=filename, content_type=part.get_content_type(), data=payload)
        )
    return parts


def parse_post_parts(headers, body: bytes) -> dict[str, list[UploadedPart]]:
    content_type = headers.get("Content-Type", "")
    if content_type.lower().startswith("application/x-www-form-urlencoded"):
        parsed = urllib.parse.parse_qs(body.decode("utf-8", errors="replace"), keep_blank_values=True)
        return {
            name: [
                UploadedPart(name=name, filename=None, content_type="text/plain", data=value.encode("utf-8"))
                for value in values
            ]
            for name, values in parsed.items()
        }
    return parse_multipart(headers, body)


def form_value(parts: dict[str, list[UploadedPart]], name: str, default: str = "") -> str:
    values = parts.get(name) or []
    if not values:
        return default
    return values[0].data.decode("utf-8", errors="replace").strip()


def file_part(parts: dict[str, list[UploadedPart]], name: str) -> UploadedPart | None:
    for part in parts.get(name) or []:
        if part.filename and part.data:
            return part
    return None


def safe_zip_extract(data: bytes, dest: Path) -> int:
    count = 0
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            raw_name = info.filename.replace("\\", "/")
            pure = PurePosixPath(raw_name)
            if pure.is_absolute() or ".." in pure.parts:
                continue
            rel = Path(*[p for p in pure.parts if p and p not in (".",)])
            target = (dest / rel).resolve()
            if not str(target).startswith(str(dest.resolve())):
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, target.open("wb") as out:
                shutil.copyfileobj(src, out, length=1024 * 1024)
            count += 1
    return count


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_path_for_match(value: str) -> str:
    value = (value or "").replace("\\", "/")
    value = re.sub(r"^[A-Za-z]:", "", value)
    value = value.strip("/")
    return posixpath.normpath(value).lower()


def build_source_index(root: Path) -> tuple[dict[str, Path], dict[str, list[Path]], dict[str, Path]]:
    suffix_map: dict[str, Path] = {}
    basename_map: dict[str, list[Path]] = {}
    sha_map: dict[str, Path] = {}
    if not root.exists():
        return suffix_map, basename_map, sha_map
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        rel = path.relative_to(root).as_posix()
        norm = norm_path_for_match(rel)
        suffix_map[norm] = path
        basename_map.setdefault(path.name.lower(), []).append(path)
        try:
            sha_map[file_sha256(path).lower()] = path
        except Exception:
            pass
    return suffix_map, basename_map, sha_map


def locate_source_file(file_path: str, sha256: str | None, source_root: Path) -> Path | None:
    suffix_map, basename_map, sha_map = build_source_index(source_root)
    sha = (sha256 or "").strip().lower()
    if sha and sha != "n/a" and sha in sha_map:
        return sha_map[sha]
    wanted = norm_path_for_match(file_path)
    for rel, path in suffix_map.items():
        if rel == wanted or wanted.endswith(rel) or rel.endswith(wanted):
            return path
    candidates = basename_map.get(Path(file_path.replace("\\", "/")).name.lower(), [])
    if len(candidates) == 1:
        return candidates[0]
    return None


def yara_externals(file_path: str, data: bytes) -> dict:
    name = Path(file_path.replace("\\", "/")).name
    ext = Path(name).suffix
    return {
        "filename": name,
        "filepath": file_path,
        "extension": ext,
        "filetype": "",
        "md5": hashlib.md5(data).hexdigest(),
        "owner": "dummy",
    }


_COMPILED_YARA = None
_YARA_STATUS = ""


def compile_yara_rules():
    global _COMPILED_YARA, _YARA_STATUS
    if _COMPILED_YARA is not None or _YARA_STATUS:
        return _COMPILED_YARA
    if yara is None:
        _YARA_STATUS = "Thiếu yara-python nên không chạy lại YARA được."
        return None

    policy_data = scanner_policy()
    rule_mode = str(policy_data.get("yara_rule_mode", "all")).lower()
    patterns = [str(p).lower() for p in policy_data.get("webshell_yara_file_patterns", [])]
    rule_dirs = [
        SCANNER_DIR / "libs" / "signature-base" / "custom",
        SCANNER_DIR / "libs" / "signature-base" / "yara",
        SCANNER_DIR / "libs" / "signature-base" / "iocs" / "yara",
    ]
    externals = {
        "filename": "",
        "filepath": "",
        "extension": "",
        "filetype": "",
        "md5": "",
        "owner": "dummy",
    }
    valid_sources: list[str] = []
    skipped = 0
    for rule_dir in rule_dirs:
        if not rule_dir.exists():
            continue
        for path in rule_dir.rglob("*"):
            if not path.is_file() or path.suffix.lower() not in (".yar", ".yara"):
                continue
            if path.name.startswith((".", "~", "_")):
                continue
            rel = path.relative_to(SCANNER_DIR).as_posix().lower()
            if rule_mode in ("webshell-only", "webshell_only", "webshell"):
                if "/custom/" not in rel and not any(p in rel for p in patterns):
                    continue
            source = path.read_text(encoding="utf-8", errors="replace")
            if not source.strip():
                continue
            try:
                yara.compile(source=source, externals=externals)
                valid_sources.append(source)
            except Exception:
                skipped += 1
    if not valid_sources:
        _YARA_STATUS = "Không compile được YARA rule nào."
        return None
    try:
        _COMPILED_YARA = yara.compile(source="\n".join(valid_sources), externals=externals)
        _YARA_STATUS = f"YARA ready: {len(valid_sources)} file rule hợp lệ, bỏ qua {skipped} file lỗi."
    except Exception as exc:
        _YARA_STATUS = f"Lỗi compile YARA tổng: {exc}"
        _COMPILED_YARA = None
    return _COMPILED_YARA


def match_instances(match) -> list[tuple[int, str, bytes]]:
    out: list[tuple[int, str, bytes]] = []
    for item in getattr(match, "strings", []) or []:
        if hasattr(item, "instances"):
            ident = str(getattr(item, "identifier", ""))
            for inst in getattr(item, "instances", []) or []:
                offset = int(getattr(inst, "offset", 0))
                data = getattr(inst, "matched_data", b"")
                if isinstance(data, str):
                    data = data.encode("utf-8", errors="replace")
                out.append((offset, ident, bytes(data)))
        elif isinstance(item, tuple) and len(item) >= 3:
            offset, ident, data = item[:3]
            if isinstance(data, str):
                data = data.encode("utf-8", errors="replace")
            out.append((int(offset), str(ident), bytes(data)))
    return out


def write_layer(case_id: int, finding_id: int, layer_name: str, data: bytes) -> str:
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", layer_name)[:80] or "layer"
    rel = Path("cases") / str(case_id) / "layers" / str(finding_id) / f"{safe_name}.bin"
    path = DATA_DIR / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return rel.as_posix()


def extract_evidence(conn: sqlite3.Connection, case_id: int, finding_id: int, file_path: str, local_file: Path) -> int:
    rules = compile_yara_rules()
    if rules is None:
        return 0
    data = local_file.read_bytes()
    policy_data = scanner_policy()
    ext = Path(file_path.replace("\\", "/")).suffix
    if iter_deobfuscated_layers and should_deobfuscate and should_deobfuscate(ext, file_path, policy_data):
        layers = list(iter_deobfuscated_layers(data, policy_data))
    else:
        layers = [("raw", data)]

    total = 0
    for layer_name, layer_data in layers:
        rel = write_layer(case_id, finding_id, layer_name, layer_data)
        conn.execute(
            "INSERT OR IGNORE INTO layers(finding_id, name, relpath, size, sha256) VALUES(?,?,?,?,?)",
            (finding_id, layer_name, rel, len(layer_data), hashlib.sha256(layer_data).hexdigest()),
        )
        try:
            matches = rules.match(
                data=layer_data,
                externals=yara_externals(file_path, layer_data),
                timeout=YARA_TIMEOUT_SECONDS,
            )
        except TypeError:
            matches = rules.match(data=layer_data, externals=yara_externals(file_path, layer_data))
        except Exception:
            continue
        for match in matches:
            meta = getattr(match, "meta", {}) or {}
            score = safe_int(meta.get("score"), 70)
            description = str(meta.get("description") or meta.get("cluster") or "not set")
            reference = str(meta.get("reference") or meta.get("viz_url") or "-")
            author = str(meta.get("author") or "-")
            instances = match_instances(match)
            if not instances:
                conn.execute(
                    """
                    INSERT INTO evidence(finding_id, layer_name, rule_name, score, description, reference, author)
                    VALUES(?,?,?,?,?,?,?)
                    """,
                    (finding_id, layer_name, match.rule, score, description, reference, author),
                )
                total += 1
            for offset, ident, matched in instances[:200]:
                end = offset + len(matched)
                text = matched[:240].decode("utf-8", errors="replace")
                conn.execute(
                    """
                    INSERT INTO evidence(
                        finding_id, layer_name, rule_name, string_id, start_offset, end_offset,
                        matched_text, score, description, reference, author
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (finding_id, layer_name, match.rule, ident, offset, end, text, score, description, reference, author),
                )
                total += 1
    return total


def _read_bundle_json(zf: zipfile.ZipFile, name: str, default):
    try:
        return json.loads(zf.read(name).decode("utf-8"))
    except KeyError:
        return default


def _safe_bundle_source_path(relpath: str) -> Path | None:
    raw = (relpath or "").replace("\\", "/")
    pure = PurePosixPath(raw)
    if pure.is_absolute() or ".." in pure.parts:
        return None
    if len(pure.parts) < 2 or pure.parts[0] != "sources":
        return None
    return Path(*[part for part in pure.parts if part and part != "."])


def _extract_bundle_member(
    zf: zipfile.ZipFile,
    relpath: str,
    source_root: Path,
    source_root_resolved: Path,
    data_root: Path,
) -> str | None:
    rel = _safe_bundle_source_path(relpath)
    if rel is None:
        return None
    try:
        data = zf.read(rel.as_posix())
    except KeyError:
        return None
    target = (source_root / rel).resolve()
    try:
        target.relative_to(source_root_resolved)
    except ValueError:
        return None
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return target.relative_to(data_root).as_posix()


def import_enc_case(title: str, enc_bundle: UploadedPart, uploaded_by_user_id=None, upload_origin="admin_web") -> tuple[int, str, int, int]:
    filename = enc_bundle.filename or "review.enc"
    if Path(filename).suffix.lower() != ".enc":
        raise RuntimeError("Chỉ chấp nhận file .enc được tạo bởi webshell-scan.")
    if upload_origin not in ("admin_web", "remote_upload"):
        raise RuntimeError("Upload origin không hợp lệ.")
    if decrypt_review_bundle_bytes is None or bundle_passphrase is None:
        raise RuntimeError("Thiếu module đọc .enc trong libs/review_bundle.py.")
    passphrase = bundle_passphrase(SCANNER_DIR, create_if_missing=False)
    zip_payload = decrypt_review_bundle_bytes(enc_bundle.data, passphrase)
    if not zipfile.is_zipfile(io.BytesIO(zip_payload)):
        raise RuntimeError("Payload .enc không chứa review bundle hợp lệ.")
    share_key = secrets.token_urlsafe(32)
    viewer_key = secrets.token_urlsafe(24)
    created = now_ts()
    with zipfile.ZipFile(io.BytesIO(zip_payload)) as zf, connect_db() as conn:
        manifest = _read_bundle_json(zf, "manifest.json", {})
        findings = _read_bundle_json(zf, "findings.json", [])
        gray_zone = _read_bundle_json(zf, "gray_zone.json", [])
        source_index = _read_bundle_json(zf, "sources/index.json", [])
        if manifest.get("schema") != "webshell-review-bundle-v1":
            raise RuntimeError("File .enc không đúng schema review bundle.")
        if not isinstance(findings, list): findings = []
        if not isinstance(gray_zone, list): gray_zone = []
        if not isinstance(source_index, list): source_index = []
        valid_rows = [row for row in findings + gray_zone if isinstance(row, dict) and str(row.get("FILE") or row.get("file") or "").strip()]
        valid_candidates = [item for item in source_index if isinstance(item, dict) and str(item.get("candidate_id") or "").strip()]
        if not valid_rows and not valid_candidates:
            raise RuntimeError("File .enc không có finding hoặc candidate hợp lệ để import.")
        if not findings and not gray_zone:
            findings = [{"FILE": str(item.get("file") or item.get("candidate_id")), "SHA256": str(item.get("sha256") or "N/A"), "SCORE": item.get("score"), "TIME": item.get("time", ""), "DESCRIPTION": item.get("description", "candidate-only import"), "EXTENSION": item.get("extension") or Path(str(item.get("file") or "")).suffix, "candidate_id": str(item.get("candidate_id")), "candidate_kind": str(item.get("candidate_kind") or "finding")} for item in valid_candidates]
        default_title = title or Path(str(manifest.get("scan_path") or "")).name or str(manifest.get("out_file") or "") or f"Scan import {time.strftime('%Y-%m-%d %H:%M')}"
        cur = conn.execute(
            "INSERT INTO cases(title, share_key, viewer_key, created_at, source_zip_name, uploaded_by_user_id, upload_origin) VALUES(?,?,?,?,?,?,?)",
            (default_title, share_key, viewer_key, created, filename, uploaded_by_user_id, upload_origin),
        )
        case_id = int(cur.lastrowid)
        data_root = DATA_DIR.resolve(); case_dir = data_root / "cases" / str(case_id); source_root = case_dir / "source"; source_root.mkdir(parents=True, exist_ok=True); source_root_resolved = source_root.resolve()
        source_by_file = {}; source_by_sha = {}; source_by_candidate = {}; source_count = 0
        for item in source_index:
            if not isinstance(item, dict): continue
            meta = {"local_relpath":"", "snippet_only":bool(item.get("snippet_only")), "layers":[], "evidence":item.get("evidence") if isinstance(item.get("evidence"), list) else [], "candidate_id":str(item.get("candidate_id") or ""), "candidate_kind":str(item.get("candidate_kind") or "finding"), "source_status":str(item.get("source_status") or "stored"), "source_reason":str(item.get("source_reason") or item.get("omitted_reason") or ""), "archive":item.get("archive") if isinstance(item.get("archive"), dict) else {}}
            if meta["candidate_kind"] not in ("finding", "gray"): meta["candidate_kind"] = "finding"
            layers = item.get("layers") if isinstance(item.get("layers"), list) else []
            for layer in layers:
                if not isinstance(layer, dict): continue
                local_layer_rel = _extract_bundle_member(zf, str(layer.get("relpath") or ""), source_root, source_root_resolved, data_root)
                if local_layer_rel:
                    layer_meta = dict(layer); layer_meta["local_relpath"] = local_layer_rel; meta["layers"].append(layer_meta); meta["local_relpath"] = meta["local_relpath"] or local_layer_rel
            if not layers:
                local_rel = _extract_bundle_member(zf, str(item.get("relpath") or ""), source_root, source_root_resolved, data_root)
                if local_rel: meta["local_relpath"] = local_rel
            if meta["local_relpath"] or meta["layers"]: source_count += 1
            if meta["candidate_id"]: source_by_candidate[meta["candidate_id"]] = meta
            if item.get("file"): source_by_file[str(item["file"])] = meta
            sha = str(item.get("sha256") or "").lower()
            if sha and sha != "n/a": source_by_sha[sha] = meta
        imported = 0; rows = [(row, str(row.get("candidate_kind") or "finding")) for row in findings if isinstance(row, dict)]; rows.extend((row, "gray") for row in gray_zone if isinstance(row, dict)); seen = set()
        for row, candidate_kind in rows:
            file_path = str(row.get("FILE") or row.get("file") or "").strip()
            if not file_path: continue
            if candidate_kind not in ("finding", "gray"): candidate_kind = "finding"
            sha256 = str(row.get("SHA256") or row.get("sha256") or "N/A").strip(); candidate_id = str(row.get("candidate_id") or ""); source_meta = source_by_candidate.get(candidate_id) or source_by_file.get(file_path) or source_by_sha.get(sha256.lower()) or {}; dedupe_key = candidate_id or "%s:%s:%s" % (candidate_kind, file_path, sha256)
            if dedupe_key in seen: continue
            seen.add(dedupe_key); local_rel = source_meta.get("local_relpath") if isinstance(source_meta, dict) else ""; local_file = DATA_DIR / local_rel if local_rel else None; bundle_layers = source_meta.get("layers", []) if isinstance(source_meta, dict) else []; bundle_evidence = source_meta.get("evidence", []) if isinstance(source_meta, dict) else []
            cur = conn.execute("INSERT INTO findings(case_id,file_path,sha256,score,scan_time,description,extension,candidate_kind,local_relpath,has_source,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)", (case_id,file_path,sha256,safe_int(row.get("SCORE",row.get("score")),None),str(row.get("TIME",row.get("time",""))),str(row.get("DESCRIPTION",row.get("description",""))),str(row.get("EXTENSION",row.get("extension",Path(file_path).suffix))),candidate_kind,local_rel,1 if (local_file and local_file.exists()) or bundle_layers else 0,created))
            finding_id = int(cur.lastrowid)
            if bundle_layers:
                stored_layer_names=set()
                for layer in bundle_layers:
                    name=str(layer.get("name") or "snippet"); rel=str(layer.get("local_relpath") or ""); path=DATA_DIR/rel
                    if not rel or not path.exists(): continue
                    stored_layer_names.add(name); conn.execute("INSERT OR IGNORE INTO layers(finding_id,name,relpath,size,sha256) VALUES(?,?,?,?,?)",(finding_id,name,rel,safe_int(layer.get("size"),path.stat().st_size),str(layer.get("sha256") or file_sha256(path))))
                for ev in bundle_evidence:
                    if not isinstance(ev,dict) or str(ev.get("layer_name") or "") not in stored_layer_names: continue
                    conn.execute("INSERT INTO evidence(finding_id,layer_name,rule_name,string_id,start_offset,end_offset,matched_text,score,description,reference,author) VALUES(?,?,?,?,?,?,?,?,?,?,?)",(finding_id,str(ev.get("layer_name") or ""),str(ev.get("rule_name") or ""),str(ev.get("string_id") or ""),safe_int(ev.get("start_offset"),None) if ev.get("start_offset") is not None else None,safe_int(ev.get("end_offset"),None) if ev.get("end_offset") is not None else None,str(ev.get("matched_text") or ""),safe_int(ev.get("score"),70),str(ev.get("description") or ""),str(ev.get("reference") or "-"),str(ev.get("author") or "-")))
            elif local_file and local_file.exists():
                try: extract_evidence(conn,case_id,finding_id,file_path,local_file)
                except Exception as exc: conn.execute("INSERT INTO audit_log(case_id,finding_id,action,detail,created_at) VALUES(?,?,?,?,?)",(case_id,finding_id,"evidence_error",str(exc),now_ts()))
            imported += 1
        conn.execute("UPDATE cases SET finding_count=?,source_count=? WHERE id=?",(imported,source_count,case_id)); conn.execute("INSERT INTO audit_log(case_id,action,detail,created_at) VALUES(?,?,?,?)",(case_id,"import_enc","Imported %s findings, %s source files from %s origin=%s"%(imported,source_count,filename,upload_origin),now_ts()))
    return case_id, share_key, imported, source_count


def import_case(title: str, findings: UploadedPart, source_zip: UploadedPart | None) -> tuple[int, str, int, int]:
    rows = parse_findings_file(findings.filename or "findings.csv", findings.data)
    if not rows:
        raise RuntimeError("Không thấy dòng finding hợp lệ trong file kết quả.")

    share_key = secrets.token_urlsafe(32)
    viewer_key = secrets.token_urlsafe(24)
    created = now_ts()
    with connect_db() as conn:
        cur = conn.execute(
            "INSERT INTO cases(title, share_key, viewer_key, created_at, source_zip_name) VALUES(?,?,?,?,?)",
            (title or f"Scan import {time.strftime('%Y-%m-%d %H:%M')}", share_key, viewer_key, created,
             source_zip.filename if source_zip else None),
        )
        case_id = int(cur.lastrowid)
        case_dir = DATA_DIR / "cases" / str(case_id)
        source_root = case_dir / "source"
        source_count = 0
        if source_zip:
            source_root.mkdir(parents=True, exist_ok=True)
            source_count = safe_zip_extract(source_zip.data, source_root)

        imported = 0
        for row in rows:
            file_path = str(row.get("FILE") or "").strip()
            if not file_path:
                continue
            sha256 = str(row.get("SHA256") or "N/A").strip()
            score = safe_int(row.get("SCORE"), None)
            local_file = locate_source_file(file_path, sha256, source_root) if source_count else None
            local_rel = local_file.relative_to(DATA_DIR).as_posix() if local_file else None
            cur = conn.execute(
                """
                INSERT INTO findings(
                    case_id, file_path, sha256, score, scan_time, description, extension,
                    local_relpath, has_source, created_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    case_id,
                    file_path,
                    sha256,
                    score,
                    str(row.get("TIME") or ""),
                    str(row.get("DESCRIPTION") or ""),
                    str(row.get("EXTENSION") or Path(file_path).suffix),
                    local_rel,
                    1 if local_file else 0,
                    created,
                ),
            )
            finding_id = int(cur.lastrowid)
            if local_file:
                try:
                    extract_evidence(conn, case_id, finding_id, file_path, local_file)
                except Exception as exc:
                    conn.execute(
                        "INSERT INTO audit_log(case_id, finding_id, action, detail, created_at) VALUES(?,?,?,?,?)",
                        (case_id, finding_id, "evidence_error", str(exc), now_ts()),
                    )
            imported += 1
        conn.execute(
            "UPDATE cases SET finding_count=?, source_count=? WHERE id=?",
            (imported, source_count, case_id),
        )
        conn.execute(
            "INSERT INTO audit_log(case_id, action, detail, created_at) VALUES(?,?,?,?)",
            (case_id, "import", f"Imported {imported} findings, {source_count} source files", now_ts()),
        )
    return case_id, share_key, imported, source_count


def query_one(sql: str, params=()):
    with connect_db() as conn:
        return conn.execute(sql, params).fetchone()


def query_all(sql: str, params=()):
    with connect_db() as conn:
        return conn.execute(sql, params).fetchall()


def score_class(score) -> str:
    score = safe_int(score, 0)
    if score >= 100:
        return "score-high"
    if score >= 40:
        return "score-mid"
    return "score-low"


def status_label(status: str) -> str:
    labels = {
        "pending": "Chờ đánh giá",
        "reviewing": "Đang xem",
        "confirmed": "Xác nhận webshell",
        "benign": "Loại trừ",
        "remediated": "Đã xử lý",
    }
    return labels.get(status, status)



def dashboard_filter_url(q: str = "", kind: str = "", status: str = "") -> str:
    values = {}
    if q:
        values["q"] = q
    if kind in ("finding", "gray"):
        values["kind"] = kind
    if status in VALID_STATUSES:
        values["status"] = status
    encoded = urllib.parse.urlencode(values)
    return "?%s" % encoded if encoded else "?"


def case_from_id(case_id: int):
    return query_one("SELECT * FROM cases WHERE id=?", (case_id,))


def case_from_key(key: str):
    return query_one("SELECT * FROM cases WHERE share_key=?", (key,))


def case_access_from_key(key: str):
    case = query_one("SELECT *, 'reviewer' AS access_role FROM cases WHERE share_key=?", (key,))
    if case:
        return case
    return query_one("SELECT *, 'viewer' AS access_role FROM cases WHERE viewer_key=?", (key,))


def create_session(case_id: int, role: str) -> str:
    token = secrets.token_urlsafe(32)
    created = now_ts()
    if role not in ("reviewer", "viewer"):
        role = "viewer"
    with connect_db() as conn:
        conn.execute("DELETE FROM sessions WHERE expires_at < ?", (created,))
        conn.execute(
            "INSERT INTO sessions(token, case_id, role, created_at, expires_at) VALUES(?,?,?,?,?)",
            (token, case_id, role, created, created + SESSION_TTL_SECONDS),
        )
    return token


def create_admin_session(user_id: int | None = None) -> tuple[str, str]:
    token = secrets.token_urlsafe(32)
    csrf_token = secrets.token_urlsafe(24)
    created = now_ts()
    with connect_db() as conn:
        conn.execute("DELETE FROM admin_sessions WHERE expires_at < ?", (created,))
        conn.execute(
            "INSERT INTO admin_sessions(token, csrf_token, user_id, created_at, expires_at) VALUES(?,?,?,?,?)",
            (token, csrf_token, user_id, created, created + SESSION_TTL_SECONDS),
        )
    return token, csrf_token


def admin_context(headers):
    token = parse_cookie(headers.get("Cookie")).get(ADMIN_COOKIE, "")
    if not token:
        return None
    row = query_one(
        "SELECT s.token, s.csrf_token, s.user_id, u.username "
        "FROM admin_sessions s LEFT JOIN users u ON u.id=s.user_id "
        "WHERE s.token=? AND s.expires_at>=?",
        (token, now_ts()),
    )
    if not row:
        return None
    return {"role": "admin", "user_id": row["user_id"], "username": row["username"] or "legacy-admin", "session_token": row["token"], "csrf_token": row["csrf_token"]}


def parse_cookie(header_value: str | None) -> dict[str, str]:
    cookie = SimpleCookie()
    try:
        cookie.load(header_value or "")
    except Exception:
        return {}
    return {key: morsel.value for key, morsel in cookie.items()}


def session_case_id(headers) -> int | None:
    info = session_info(headers)
    return info[0] if info else None


def revoke_sessions(headers) -> None:
    tokens = {SESSION_COOKIE: set(), ADMIN_COOKIE: set(), USER_COOKIE: set()}
    for part in (headers.get("Cookie") or "").split(";"):
        name, separator, value = part.strip().partition("=")
        if separator and name in tokens and value:
            tokens[name].add(value.strip())
    with connect_db() as conn:
        for table, cookie_name in (("sessions", SESSION_COOKIE), ("admin_sessions", ADMIN_COOKIE), ("user_sessions", USER_COOKIE)):
            for token in tokens[cookie_name]:
                conn.execute(f"DELETE FROM {table} WHERE token=?", (token,))

def session_info(headers) -> tuple[int, str] | None:
    token = parse_cookie(headers.get("Cookie")).get(SESSION_COOKIE, "")
    if not token:
        return None
    now = now_ts()
    row = query_one("SELECT case_id, role FROM sessions WHERE token=? AND expires_at>=?", (token, now))
    if not row:
        return None
    return int(row["case_id"]), str(row["role"] or "viewer")



def create_user_session(user_id: int) -> tuple[str, str]:
    token = secrets.token_urlsafe(32)
    csrf_token = secrets.token_urlsafe(24)
    created = now_ts()
    with connect_db() as conn:
        conn.execute("DELETE FROM user_sessions WHERE expires_at < ?", (created,))
        conn.execute(
            "INSERT INTO user_sessions(token,user_id,csrf_token,created_at,expires_at) VALUES(?,?,?,?,?)",
            (token, user_id, csrf_token, created, created + SESSION_TTL_SECONDS),
        )
    return token, csrf_token


def user_context(headers):
    token = parse_cookie(headers.get("Cookie")).get("review_user", "")
    if not token:
        return None
    row = query_one(
        "SELECT s.token, s.user_id, s.csrf_token, u.username, u.role, u.active "
        "FROM user_sessions s JOIN users u ON u.id=s.user_id "
        "WHERE s.token=? AND s.expires_at>=?",
        (token, now_ts()),
    )
    if not row or not row["active"]:
        return None
    return dict(row)


def csrf_valid(context, supplied: str) -> bool:
    return bool(context and supplied and hmac.compare_digest(str(context.get("csrf_token", "")), str(supplied)))


def report_permission(case, context, required="view") -> bool:
    if not context:
        return False
    if context.get("role") == "admin":
        return True
    if context.get("role") != "user" or required not in ("view", "process"):
        return False
    return query_one(
        "SELECT 1 FROM case_assignments WHERE case_id=? AND user_id=? AND active=1",
        (case["id"], context["user_id"]),
    ) is not None


def case_visible_to_user(case_id: int, user_id: int) -> bool:
    return query_one(
        "SELECT 1 FROM case_assignments WHERE case_id=? AND user_id=? AND active=1",
        (case_id, user_id),
    ) is not None



def assign_report(case_id: int, user_id: int, admin_user_id: int) -> None:
    user = query_one("SELECT id FROM users WHERE id=? AND role='user' AND active=1", (user_id,))
    if not user:
        raise ValueError("Only active user accounts can be assigned")
    with connect_db() as conn:
        existing = conn.execute("SELECT id FROM case_assignments WHERE case_id=? AND user_id=? AND active=1", (case_id, user_id)).fetchone()
        if not existing:
            conn.execute("INSERT INTO case_assignments(case_id,user_id,active,assigned_by_user_id,created_at) VALUES(?,?,?,?,?)", (case_id, user_id, 1, admin_user_id, now_ts()))
        conn.execute("INSERT INTO audit_log(case_id,action,detail,created_at) VALUES(?,?,?,?)", (case_id, "report_assign", "actor=%s user=%s" % (admin_user_id, user_id), now_ts()))


def revoke_report(case_id: int, user_id: int, admin_user_id: int) -> None:
    with connect_db() as conn:
        conn.execute("UPDATE case_assignments SET active=0, revoked_at=? WHERE case_id=? AND user_id=? AND active=1", (now_ts(), case_id, user_id))
        conn.execute("INSERT INTO audit_log(case_id,action,detail,created_at) VALUES(?,?,?,?)", (case_id, "report_revoke", "actor=%s user=%s" % (admin_user_id, user_id), now_ts()))


def assigned_users(case_id: int):
    return query_all("SELECT u.id,u.username FROM case_assignments a JOIN users u ON u.id=a.user_id WHERE a.case_id=? AND a.active=1 AND u.active=1 ORDER BY u.username", (case_id,))



def rotate_upload_key(actor_user_id: int) -> str:
    value = secrets.token_urlsafe(48)
    now = now_ts()
    with connect_db() as conn:
        conn.execute("UPDATE upload_credentials SET active=0, rotated_at=? WHERE active=1", (now,))
        conn.execute(
            "INSERT INTO upload_credentials(key_hash,active,created_at,created_by_user_id) VALUES(?,?,?,?)",
            (hash_password(value), 1, now, actor_user_id),
        )
        conn.execute(
            "INSERT INTO audit_log(action,detail,created_at) VALUES(?,?,?)",
            ("upload_key_rotate", "actor=%s" % actor_user_id, now),
        )
    return value


def verify_upload_key(value: str) -> bool:
    if not value:
        return False
    rows = query_all("SELECT key_hash FROM upload_credentials WHERE active=1 ORDER BY id DESC LIMIT 1")
    return bool(rows and verify_password(value, rows[0]["key_hash"]))



def update_finding_status(case_id: int, finding_id: int, status: str, note: str, context: dict) -> None:
    if status not in VALID_STATUSES:
        raise ValueError("Invalid status")
    if status == "benign" and not str(note or "").strip():
        raise ValueError("Benign status requires a note")
    case = case_from_id(case_id)
    if not case or not report_permission(case, context, "process"):
        raise PermissionError("Report access denied")
    with connect_db() as conn:
        row = conn.execute("SELECT id FROM findings WHERE id=? AND case_id=?", (finding_id, case_id)).fetchone()
        if not row:
            raise PermissionError("Finding access denied")
        conn.execute("UPDATE findings SET status=?, reviewer_note=? WHERE id=? AND case_id=?", (status, note, finding_id, case_id))
        conn.execute("INSERT INTO audit_log(case_id,finding_id,action,detail,created_at) VALUES(?,?,?,?,?)", (case_id, finding_id, "review_update", "note_present=%s" % bool(str(note or "").strip()), now_ts()))


def bulk_update_findings(case_id: int, finding_ids: list[int], status: str, note: str, context: dict) -> int:
    if not finding_ids:
        return 0
    if len(finding_ids) > 500:
        raise ValueError("Too many selected findings")
    if status not in ("reviewing", "confirmed", "remediated", "benign"):
        raise ValueError("Invalid bulk status")
    if status == "benign" and not str(note or "").strip():
        raise ValueError("Benign status requires a note")
    case = case_from_id(case_id)
    if not case or not report_permission(case, context, "process"):
        raise PermissionError("Report access denied")
    ids = sorted({int(value) for value in finding_ids})
    placeholders = ",".join("?" for _ in ids)
    with connect_db() as conn:
        rows = conn.execute("SELECT id FROM findings WHERE case_id=? AND id IN (%s)" % placeholders, [case_id, *ids]).fetchall()
        if len(rows) != len(ids):
            raise ValueError("Selection contains an invalid finding")
        conn.execute("UPDATE findings SET status=?, reviewer_note=? WHERE case_id=? AND id IN (%s)" % placeholders, [status, note, case_id, *ids])
        conn.execute("INSERT INTO audit_log(case_id,action,detail,created_at) VALUES(?,?,?,?)", (case_id, "bulk_review_update", "status=%s selected=%s note_present=%s" % (status, len(ids), bool(str(note or "").strip())), now_ts()))
    return len(ids)



def create_customer_user(username: str, password: str, actor_user_id: int) -> int:
    normalized = normalize_username(username)
    if not normalized or normalized == "admin":
        raise ValueError("Invalid username")
    now = now_ts()
    with connect_db() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO users(username,password_hash,role,active,must_change_password,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
                (normalized, hash_password(password), "user", 1, 0, now, now),
            )
        except sqlite3.IntegrityError as exc:
            raise ValueError("Username already exists") from exc
        conn.execute(
            "INSERT INTO audit_log(action,detail,created_at) VALUES(?,?,?)",
            ("user_create", "actor=%s user=%s" % (actor_user_id, normalized), now),
        )
        return int(cur.lastrowid)


def set_user_active(user_id: int, active: bool, actor_user_id: int) -> None:
    with connect_db() as conn:
        conn.execute("UPDATE users SET active=?, updated_at=? WHERE id=? AND role='user'", (int(active), now_ts(), user_id))
        if not active:
            conn.execute("DELETE FROM user_sessions WHERE user_id=?", (user_id,))
        conn.execute(
            "INSERT INTO audit_log(action,detail,created_at) VALUES(?,?,?)",
            ("user_active" if active else "user_disabled", "actor=%s user=%s" % (actor_user_id, user_id), now_ts()),
        )


def reset_user_password(user_id: int, password: str, actor_user_id: int) -> None:
    with connect_db() as conn:
        conn.execute(
            "UPDATE users SET password_hash=?, must_change_password=1, updated_at=? WHERE id=? AND role='user'",
            (hash_password(password), now_ts(), user_id),
        )
        conn.execute("DELETE FROM user_sessions WHERE user_id=?", (user_id,))
        conn.execute(
            "INSERT INTO audit_log(action,detail,created_at) VALUES(?,?,?)",
            ("user_password_reset", "actor=%s user=%s" % (actor_user_id, user_id), now_ts()),
        )


def is_admin_session(headers) -> bool:
    token = parse_cookie(headers.get("Cookie")).get(ADMIN_COOKIE, "")
    if not token:
        return False
    now = now_ts()
    return query_one("SELECT token FROM admin_sessions WHERE token=? AND expires_at>=?", (token, now)) is not None


def admin_login_page(message: str = "") -> bytes:
    notice = f"<div class=\"warn\">{html.escape(message)}</div>" if message else ""
    body = f"""
    <section class="band" style="max-width:560px;margin:42px auto;">
      <h1>Đăng nhập trang import</h1>
      {notice}
      <form method="post" action="/admin/login" class="stack">
        <label>Admin key<br><input name="admin_key" type="password" autocomplete="current-password" required style="width:100%"></label>
        <button type="submit">Vào trang admin</button>
      </form>
    </section>
    """
    return html_page("Admin login", body)


def access_page(case_id: int | None = None, message: str = "") -> bytes:
    hidden = f'<input type="hidden" name="case_id" value="{case_id}">' if case_id else ""
    notice = f"<div class=\"warn\">{html.escape(message)}</div>" if message else ""
    body = f"""
    <section class="band" style="max-width:560px;margin:42px auto;">
      <h1>Nhập mã truy cập báo cáo</h1>
      {notice}
      <form method="post" action="/access" class="stack">
        {hidden}
        <label>Mã truy cập<br><input name="access_key" type="password" autocomplete="current-password" required style="width:100%"></label>
        <button type="submit">Xem dashboard</button>
      </form>
    </section>
    """
    return html_page("Nhập mã truy cập", body)


def admin_page_for_context(base_url: str, message: str = "", admin_context_value=None) -> bytes:
    cases = query_all("SELECT * FROM cases ORDER BY created_at DESC LIMIT 100")
    users = query_all("SELECT u.id,u.username,u.active,(SELECT COUNT(*) FROM case_assignments a WHERE a.user_id=u.id AND a.active=1) AS assigned_count FROM users u WHERE u.role='user' ORDER BY u.username")
    csrf = html.escape((admin_context_value or {}).get("csrf_token", "")); username=html.escape((admin_context_value or {}).get("username", "admin"))
    active_users=sum(bool(u["active"]) for u in users); unassigned=0; main_total=gray_total=0; rows=[]
    user_options="".join(f"<option value='{u['id']}'>{html.escape(u['username'])}</option>" for u in users if u["active"])
    for case in cases:
        kind_counts={row["candidate_kind"]:row["c"] for row in query_all("SELECT candidate_kind,COUNT(*) AS c FROM findings WHERE case_id=? GROUP BY candidate_kind",(case["id"],))}; main_total+=kind_counts.get("finding",0); gray_total+=kind_counts.get("gray",0)
        assignments=query_all("SELECT u.id,u.username FROM case_assignments a JOIN users u ON u.id=a.user_id WHERE a.case_id=? AND a.active=1 ORDER BY u.username",(case["id"],)); assigned=", ".join(html.escape(row["username"]) for row in assignments) or "Chưa gán"; unassigned += not bool(assignments)
        assign_form=f"<form method='post' action='/admin/cases/{case['id']}/assign'><input type='hidden' name='csrf_token' value='{csrf}'><select name='user_id' required><option value=''>Gán user</option>{user_options}</select><button class='secondary' type='submit'>Gán</button></form>"
        revoke_forms="".join(f"<form method='post' action='/admin/cases/{case['id']}/revoke'><input type='hidden' name='csrf_token' value='{csrf}'><input type='hidden' name='user_id' value='{row['id']}'><button class='secondary' type='submit'>Bỏ {html.escape(row['username'])}</button></form>" for row in assignments)
        rows.append(f"<tr><td data-label='Tên'><a href='/case/{case['id']}'>{html.escape(case['title'])}</a></td><td data-label='Origin'>{html.escape(case['upload_origin'] or 'admin_web')}</td><td data-label='Main'><span class='kind-finding'>{kind_counts.get('finding',0)}</span></td><td data-label='Gray'><span class='kind-gray'>{kind_counts.get('gray',0)}</span></td><td data-label='Người xử lý'><span class='assignment-summary'>{assigned}</span><details class='assignment-panel'><summary>Assignment</summary>{assign_form}{revoke_forms}</details></td><td data-label='Action'><a class='button secondary' href='/case/{case['id']}'>Open</a></td></tr>")
    user_rows="".join(f"<tr><td data-label='Username'>{html.escape(u['username'])}</td><td data-label='Trạng thái'>{'Đang hoạt động' if u['active'] else 'Đã vô hiệu hóa'}</td><td data-label='Được gán'>{u['assigned_count']}</td><td data-label='Action'><form method='post' action='/admin/users/{u['id']}/active'><input type='hidden' name='csrf_token' value='{csrf}'><input type='hidden' name='active' value='{0 if u['active'] else 1}'><button class='secondary' type='submit'>{'Kích hoạt' if not u['active'] else 'Vô hiệu hóa'}</button></form></td></tr>" for u in users)
    notice=f"<div class='notice'>{html.escape(message)}</div>" if message else ""
    body=f"""<!-- soc-admin -->
    <section class="admin-header"><div><div class="eyebrow">STAGING · Admin control plane</div><h1>Review workspace</h1><p class="muted">Reports, assignments and user access.</p>{notice}</div><div class="admin-identity"><b>{username}</b><span class="pill">ADMIN</span><a class="button secondary" href="/admin/logout">Logout</a></div></section>
    <section class="action-bar"><a class="button" href="#import">Import .enc</a><a class="button secondary" href="#users">Create user</a><a class="button secondary sensitive-action" href="#upload-key">Manage upload key</a><a class="button secondary" href="#audit">Audit log</a></section>
    <section class="kpi-row"><div class="kpi"><b>{len(cases)}</b><span>Reports</span></div><div class="kpi"><b>{unassigned}</b><span>Chưa gán</span></div><div class="kpi"><b>{active_users}</b><span>Active users</span></div><div class="kpi"><b>{main_total+gray_total}</b><span>Main + gray</span></div></section>
    <section class="band upload-card" id="import"><h2>Import .enc</h2><form method="post" action="/admin/import" enctype="multipart/form-data" class="stack"><input type="hidden" name="csrf_token" value="{csrf}"><label>Tên đợt rà quét<br><input name="title" style="width:100%"></label><label>File kết quả<br><input type="file" name="review_bundle" required accept=".enc"></label><button type="submit">Import bundle</button></form></section>
    <section class="band" id="users"><h2>Users</h2><form method="post" action="/admin/users/create" class="filters"><input type="hidden" name="csrf_token" value="{csrf}"><input name="username" placeholder="username" required><input name="password" type="password" placeholder="Mật khẩu >=12 ký tự" required><button type="submit">Create user</button></form><div class="tablewrap"><table><thead><tr><th>Username</th><th>Trạng thái</th><th>Được gán</th><th>Action</th></tr></thead><tbody>{user_rows or '<tr><td colspan="4">Chưa có user.</td></tr>'}</tbody></table></div></section>
    <section class="band"><h2>Reports</h2><div class="tablewrap"><table class="report-table"><thead><tr><th>Tên</th><th>Origin</th><th>Main</th><th>Gray</th><th>Người xử lý</th><th>Action</th></tr></thead><tbody>{''.join(rows) or '<tr><td colspan="6">Chưa có báo cáo.</td></tr>'}</tbody></table></div></section>
    <section class="band" id="upload-key"><div class="warn">Upload key và audit controls chỉ dành cho admin. Không hiển thị secret trên giao diện.</div><span id="audit"></span></section>
    """
    return html_page("Admin control plane",body)


def case_dashboard(case_id: int, query: dict[str, list[str]], context=None) -> bytes:
    case = case_from_id(case_id)
    if not case:
        return html_page("Không tìm thấy", "<section class=\"band\">Báo cáo không tồn tại.</section>")
    status = (query.get("status") or [""])[0]
    kind = (query.get("kind") or [""])[0]
    q = (query.get("q") or [""])[0].strip()
    where = ["case_id=?"]
    params = [case["id"]]
    if status in VALID_STATUSES:
        where.append("status=?")
        params.append(status)
    if kind in ("finding", "gray"):
        where.append("candidate_kind=?")
        params.append(kind)
    if q:
        where.append("(file_path LIKE ? OR sha256 LIKE ? OR description LIKE ?)")
        like = f"%{q}%"
        params.extend([like, like, like])
    findings = query_all(
        "SELECT f.*,(SELECT COUNT(*) FROM evidence e WHERE e.finding_id=f.id) AS evidence_count "
        "FROM findings f WHERE %s ORDER BY score DESC,id ASC LIMIT 1000" % " AND ".join(where),
        params,
    )
    counts = {row["status"]: row["c"] for row in query_all("SELECT status,COUNT(*) AS c FROM findings WHERE case_id=? GROUP BY status", (case["id"],))}
    kind_counts = {row["candidate_kind"]: row["c"] for row in query_all("SELECT candidate_kind,COUNT(*) AS c FROM findings WHERE case_id=? GROUP BY candidate_kind", (case["id"],))}
    admin = bool(context and context.get("role") == "admin")
    csrf = html.escape((context or {}).get("csrf_token", ""))
    process = admin or bool(context and report_permission(case, context, "process"))
    metric_html = """
    <div class="grid">
      <div class="metric metric-main"><b>%s</b><span>Phát hiện chính</span></div>
      <div class="metric metric-gray"><b>%s</b><span>Gray zone</span></div>
      <div class="metric metric-ok"><b>%s</b><span>Đã xác nhận</span></div>
      <div class="metric"><b>%s</b><span>Source files</span></div>
    </div>
    """ % (kind_counts.get("finding", 0), kind_counts.get("gray", 0), counts.get("confirmed", 0), case["source_count"])
    filters = [
        "<a class='pill %s' href='%s'>Tất cả</a>" % ("filter-active" if not kind else "", dashboard_filter_url(q, "", status)),
        "<a class='pill %s' href='%s'>Phát hiện chính: %s</a>" % ("filter-active" if kind == "finding" else "", dashboard_filter_url(q, "finding", status), kind_counts.get("finding", 0)),
        "<a class='pill %s' href='%s'>Gray zone: %s</a>" % ("filter-active" if kind == "gray" else "", dashboard_filter_url(q, "gray", status), kind_counts.get("gray", 0)),
    ]
    for st in VALID_STATUSES:
        filters.append("<a class='pill status-%s %s' href='%s'>%s: %s</a>" % (st, "filter-active" if status == st else "", dashboard_filter_url(q, kind, st), status_label(st), counts.get(st, 0)))
    rows = []
    for item in findings:
        kind_class = "kind-gray" if item["candidate_kind"] == "gray" else "kind-finding"
        kind_label = "Gray zone" if item["candidate_kind"] == "gray" else "Phát hiện chính"
        source_class = "source-ok" if item["has_source"] else "source-missing"
        source_label = "Có source" if item["has_source"] else "Thiếu source"
        check = '<input type="checkbox" name="finding_ids" value="%s" form="bulk-form">' % item["id"] if process else ""
        rows.append(
            "<tr class='table-card'><td data-label='Chọn'>%s</td><td data-label='ID'><a href='/case/%s/finding/%s' class='finding-link'>#%s</a></td>"
            "<td data-label='Loại'><span class='%s'>%s</span></td><td data-label='Score' class='%s'>%s</td>"
            "<td data-label='Trạng thái'><span class='pill status-%s'>%s</span></td><td data-label='File' class='path'><a href='/case/%s/finding/%s' class='path-link'>%s</a><br><span class='muted'>%s</span></td>"
            "<td data-label='Evidence'>%s</td><td data-label='Source'><span class='%s'>%s</span></td></tr>"
            % (check, case_id, item["id"], item["id"], kind_class, kind_label, score_class(item["score"]), item["score"] or "", item["status"], status_label(item["status"]), case_id, item["id"], html.escape(item["file_path"]), html.escape(item["sha256"] or "N/A"), item["evidence_count"], source_class, source_label)
        )
    bulk = ""
    if process:
        bulk = '<form method="post" action="/case/%s/findings/bulk" id="bulk-form" class="band filters"><input type="hidden" name="csrf_token" value="%s"><select name="status"><option value="reviewing">Đang xem</option><option value="confirmed">Xác nhận webshell</option><option value="remediated">Đã xử lý</option><option value="benign">Bỏ qua</option></select><input name="note" placeholder="Ghi chú bắt buộc khi bỏ qua"><button type="submit">Cập nhật mục đã chọn</button></form>' % (case_id, csrf)
    shell_marker = "<!-- soc-admin -->" if admin else ""
    body = shell_marker + """
    <section class="band page-head hero"><div class="eyebrow">Tổng quan kết quả rà quét</div><h1>Dashboard: %s</h1><p class="muted">Main findings và gray zone được tách riêng.</p>%s</section>
    <section class="band"><form method="get" class="filters"><input name="q" value="%s" placeholder="Tìm path, hash, mô tả" style="min-width:280px"><button type="submit">Lọc</button>%s</form></section>%s
    <section class="band"><div class="tablewrap"><table><thead><tr><th>Chọn</th><th>ID</th><th>Loại</th><th>Score</th><th>Trạng thái</th><th>File</th><th>Evidence</th><th>Source</th></tr></thead><tbody>%s</tbody></table></div></section>
    """ % (html.escape(case["title"]), metric_html, html.escape(q), "".join(filters), bulk, "".join(rows) or '<tr><td colspan="8">Không có finding phù hợp.</td></tr>')
    return html_page(case["title"], body)


def read_layer_bytes(finding_id: int, layer_name: str) -> tuple[bytes | None, str]:
    row = query_one("SELECT * FROM layers WHERE finding_id=? AND name=?", (finding_id, layer_name))
    if row:
        path = DATA_DIR / row["relpath"]
        if path.exists() and path.stat().st_size <= MAX_VIEW_BYTES:
            return path.read_bytes(), row["name"]
        return None, row["name"]
    finding = query_one("SELECT * FROM findings WHERE id=?", (finding_id,))
    if finding and finding["local_relpath"]:
        path = DATA_DIR / finding["local_relpath"]
        if path.exists() and path.stat().st_size <= MAX_VIEW_BYTES:
            return path.read_bytes(), "raw"
    return None, layer_name


def render_code(data: bytes, highlights: list[sqlite3.Row]) -> str:
    spans: list[tuple[int, int, int]] = []
    for idx, ev in enumerate(highlights):
        start = ev["start_offset"]
        end = ev["end_offset"]
        if start is None or end is None:
            continue
        start_i, end_i = int(start), int(end)
        if 0 <= start_i < end_i <= len(data):
            spans.append((start_i, end_i, idx % 3))
    spans.sort(key=lambda x: (x[0], -(x[1] - x[0])))
    chunks = []
    cursor = 0
    for start, end, color in spans:
        if start < cursor:
            continue
        chunks.append(html.escape(data[cursor:start].decode("utf-8", errors="replace")))
        chunks.append(f"<mark class=\"m{color}\">")
        chunks.append(html.escape(data[start:end].decode("utf-8", errors="replace")))
        chunks.append("</mark>")
        cursor = end
    chunks.append(html.escape(data[cursor:].decode("utf-8", errors="replace")))
    rendered = "".join(chunks)
    lines = rendered.splitlines() or [""]
    numbered = []
    for idx, line in enumerate(lines, start=1):
        numbered.append(f"<span class=\"ln\">{idx}</span>{line}")
    return "\n".join(numbered)


def finding_detail(case_id: int, finding_id: int, query: dict[str, list[str]], role: str = "viewer") -> bytes:
    return _render_finding_detail(case_id, finding_id, query, role)

def _render_finding_detail(case_id: int, finding_id: int, query: dict[str, list[str]], role: str = "viewer", context=None) -> bytes:
    csrf = html.escape((context or {}).get("csrf_token", ""))
    case = case_from_id(case_id)
    if not case:
        return html_page("Không tìm thấy", "<section class=\"band\">Báo cáo không tồn tại.</section>")
    finding = query_one("SELECT * FROM findings WHERE id=? AND case_id=?", (finding_id, case["id"]))
    if not finding:
        return html_page("Không tìm thấy", "<section class=\"band\">Finding không tồn tại.</section>")
    layers = query_all(
        """
        SELECT * FROM layers
        WHERE finding_id=?
        ORDER BY CASE
            WHEN name='raw' THEN 0
            WHEN name='raw_snippet' THEN 1
            WHEN name LIKE '%_snippet' THEN 2
            ELSE 3
        END, name
        """,
        (finding_id,),
    )
    selected_layer = (query.get("layer") or ["raw"])[0]
    if layers and selected_layer not in [row["name"] for row in layers]:
        selected_layer = layers[0]["name"]
    evidence = query_all(
        "SELECT * FROM evidence WHERE finding_id=? AND layer_name=? ORDER BY rule_name, start_offset",
        (finding_id, selected_layer),
    )
    all_evidence = query_all(
        "SELECT rule_name, layer_name, COUNT(*) AS c FROM evidence WHERE finding_id=? GROUP BY rule_name, layer_name ORDER BY rule_name",
        (finding_id,),
    )
    data, actual_layer = read_layer_bytes(finding_id, selected_layer)
    kind_is_gray = str(finding["candidate_kind"] or "finding") == "gray"
    kind_badge = '<span class="kind-gray">Gray zone</span>' if kind_is_gray else '<span class="kind-finding">Phát hiện chính</span>'
    source_badge = '<span class="source-ok">Có source</span>' if finding["has_source"] else '<span class="source-missing">Thiếu source</span>'
    archive_badge = '<span class="pill">Archive member</span>' if "[ARCHIVE:" in str(finding["file_path"] or "") else ""
    layer_links = []
    for layer in layers:
        active = " filter-active" if layer["name"] == selected_layer else ""
        layer_links.append(
            f"<a class=\"pill{active}\" href=\"?layer={urllib.parse.quote(layer['name'])}\">{html.escape(layer['name'])} ({layer['size']} bytes)</a>"
        )
    rule_rows = []
    for ev in evidence:
        matched = html.escape((ev["matched_text"] or "")[:180])
        offset = "" if ev["start_offset"] is None else f"@{ev['start_offset']}"
        rule_rows.append(
            f"<div class=\"rule-hit\"><b>{html.escape(ev['rule_name'])}</b> "
            f"<span class=\"muted\">{html.escape(ev['string_id'] or '')} {offset}</span><br>"
            f"<code>{matched}</code><br><span class=\"muted\">{html.escape(ev['description'] or '')}</span></div>"
        )
    if data is None:
        code_html = '<div class="warn source-missing">Không có source/layer để hiển thị. Hãy import lại kèm source/evidence cần thiết từ scanner.</div>'
    else:
        code_html = f'<div class="codewrap"><div class="code-toolbar"><span>Raw/decode viewer: {html.escape(actual_layer)}</span><span>read-only</span></div><pre class="code">{render_code(data, evidence)}</pre></div>'
    snippet_notice = ""
    if str(actual_layer).endswith("_snippet"):
        snippet_notice = '<div class="notice snippet-only"><b>Snippet-only.</b> Hệ thống chỉ đóng gói các đoạn quanh rule hit để giữ file .enc gọn, không phải toàn bộ source.</div>'
    manual_review_notice = (
        '<div class="notice manual-review"><b>Gray zone — cần review thủ công.</b> Điểm số là tín hiệu, chưa phải kết luận webshell.</div>'
        if kind_is_gray else ""
    )
    missing_notice = (
        '<div class="warn source-missing"><b>Thiếu source.</b> Bundle không có nội dung đầy đủ để hiển thị.</div>'
        if not finding["has_source"] else ""
    )
    evidence_summary = ", ".join(
        f"{html.escape(row['rule_name'])} [{html.escape(row['layer_name'])}: {row['c']}]" for row in all_evidence[:40]
    )
    description = str(finding["description"] or "")
    description_block = ""
    if description:
        description_block = f"""
        <details class="details">
          <summary>Mô tả scanner ({len(description)} ký tự)</summary>
          <p class="muted detail-text">{html.escape(description)}</p>
        </details>
        """
    evidence_block = ""
    if evidence_summary:
        evidence_block = f"""
        <details class="details">
          <summary>Tóm tắt rule hit ({len(all_evidence)} nhóm)</summary>
          <p class="muted detail-text">{evidence_summary}</p>
        </details>
        """
    if role == "reviewer":
        review_panel = f"""
        <section class="detail-section review-controls">
          <h2>Review decision</h2>
          <form method="post" action="/case/{case['id']}/finding/{finding_id}/update" enctype="multipart/form-data" class="stack">
            <input type="hidden" name="csrf_token" value="{csrf}">
            <label>Trạng thái<br>
              <select name="status">
                {''.join(f'<option value="{st}" {"selected" if finding["status"] == st else ""}>{status_label(st)}</option>' for st in VALID_STATUSES)}
              </select>
            </label>
            <label>Ghi chú đánh giá<br><textarea name="reviewer_note">{html.escape(finding['reviewer_note'] or '')}</textarea></label>
            <button type="submit">Lưu đánh giá</button>
          </form>
        </section>
        """
    else:
        review_panel = f"""
        <section class="detail-section review-controls">
          <h2>Review decision</h2>
          <div class="notice">Bạn đang xem bằng mã lãnh đạo nên chỉ có quyền xem.</div>
          <p class="detail-text">{html.escape(finding['reviewer_note'] or 'Chưa có ghi chú.')}</p>
        </section>
        """
    file_path = html.escape(finding["file_path"] or "")
    sha256 = html.escape(finding["sha256"] or "N/A")
    detail_meta = """
    <section class="detail-meta">
      <div class="eyebrow">Finding #%s · read-only evidence</div>
      <h1 class="path">%s</h1>
      <div class="filters">%s<span class="pill %s">Score %s</span><span class="pill status-%s">%s</span>%s%s</div>
      <div class="meta-actions"><span class="path">SHA256 %s</span><button type="button" class="secondary" data-copy="sha256" data-copy-value="%s">Copy SHA256</button></div>
    </section>
    """ % (finding_id, file_path, kind_badge, score_class(finding["score"]), html.escape(str(finding["score"] or "")), finding["status"], status_label(finding["status"]), source_badge, archive_badge, sha256, sha256)
    action_html = ""
    if role == "reviewer":
        action_html = (
            '<div class="action-group" aria-label="Thao tác xử lý">'
            '<button type="button" class="action-confirm">Xác nhận</button>'
            '<button type="button" class="action-remediate">Đã xử lý</button>'
            '<button type="button" class="action-benign">Bỏ qua + ghi chú</button>'
            '</div>'
        )
    shell_marker = "<!-- soc-admin -->" if context and context.get("role") == "admin" else ""
    body = shell_marker + """
    %s%s%s%s%s
    <div class="split detail-layout"><aside class="band detail-sidebar stack"><div class="detail-section"><h2>Layer</h2><div class="filters">%s</div></div><div class="detail-section"><h2>Rule/String hit</h2><div class="rule-list">%s</div></div>%s%s</aside><section class="band detail-viewer evidence-viewer">%s</section></div>
    """ % (detail_meta, manual_review_notice, missing_notice, snippet_notice, description_block + (evidence_block or "<p class='muted'>Chưa có evidence chi tiết.</p>"), "".join(layer_links) or '<span class="muted">Chưa có layer.</span>', "".join(rule_rows) or '<div class="muted">Không có hit trên layer này.</div>', review_panel, action_html, code_html)
    return html_page(f"Finding #{finding_id}", body)


class ReviewHandler(BaseHTTPRequestHandler):
    server_version = "WebshellReview/1.0"

    def end_headers(self) -> None:
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Content-Security-Policy", "default-src 'self'; style-src 'unsafe-inline' 'self'; frame-ancestors 'none'")
        super().end_headers()

    def send_html(self, data: bytes, status: HTTPStatus = HTTPStatus.OK) -> None:
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


    def send_json(self, payload: dict, status: HTTPStatus = HTTPStatus.OK) -> None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def redirect(self, location: str, cookie: str | None = None, cookies: tuple[str, ...] = ()) -> None:
        self.send_response(HTTPStatus.SEE_OTHER)
        self.send_header("Location", location)
        if cookie:
            self.send_header("Set-Cookie", cookie)
        for value in cookies:
            self.send_header("Set-Cookie", value)
        self.end_headers()

    def read_body(self) -> bytes:
        length = int(self.headers.get("Content-Length", "0") or "0")
        if length > MAX_UPLOAD_BYTES:
            raise RuntimeError("Upload quá lớn.")
        return self.rfile.read(length)

    def base_url(self) -> str:
        host = self.headers.get("Host") or f"{self.server.server_address[0]}:{self.server.server_address[1]}"
        return f"http://{host}"

    def do_GET(self) -> None:
        parsed = urllib.parse.urlparse(self.path); query = urllib.parse.parse_qs(parsed.query)
        if parsed.path == "/static/review.js":
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "application/javascript; charset=utf-8")
            self.send_header("Content-Length", str(len(REVIEW_JS)))
            self.end_headers()
            self.wfile.write(REVIEW_JS)
            return
        if parsed.path == "/":
            context = user_context(self.headers); admin = admin_context(self.headers); self.redirect("/admin" if admin else "/user" if context else "/login"); return
        if parsed.path == "/logout":
            self.redirect("/admin/logout"); return
        if parsed.path == "/admin/logout":
            revoke_sessions(self.headers)
            expired = (
                f"{SESSION_COOKIE}=; Path=/; Max-Age=0; HttpOnly; SameSite=Strict",
                f"{ADMIN_COOKIE}=; Path=/; Max-Age=0; HttpOnly; SameSite=Strict",
                f"{ADMIN_COOKIE}=; Path=/admin; Max-Age=0; HttpOnly; SameSite=Strict",
                f"{USER_COOKIE}=; Path=/; Max-Age=0; HttpOnly; SameSite=Strict",
            )
            self.redirect("/login", cookies=expired); return
        if parsed.path == "/login": self.send_html(login_page()); return
        if parsed.path == "/admin":
            admin = admin_context(self.headers)
            if not admin: self.send_html(admin_login_page()); return
            self.send_html(admin_page_for_context(self.base_url(), admin_context_value=admin)); return
        if parsed.path == "/user":
            context = user_context(self.headers)
            if not context or context["role"] != "user": self.send_html(login_page("Cần đăng nhập tài khoản user."), HTTPStatus.FORBIDDEN); return
            cases = query_all("SELECT c.* FROM cases c JOIN case_assignments a ON a.case_id=c.id WHERE a.user_id=? AND a.active=1 ORDER BY c.created_at DESC", (context["user_id"],))
            rows = "".join(f"<tr><td><a href='/case/{row['id']}'>{html.escape(row['title'])}</a></td><td>{html.escape(row['upload_origin'] or '')}</td></tr>" for row in cases)
            empty_rows = '<tr><td colspan="2">Chưa có báo cáo được gán.</td></tr>'
            body = "<section class='band hero'><h1>Báo cáo được gán</h1><p>Chỉ báo cáo active assignment của tài khoản hiện tại.</p></section><section class='band'><table><thead><tr><th>Báo cáo</th><th>Origin</th></tr></thead><tbody>%s</tbody></table></section>" % (rows or empty_rows)
            self.send_html(html_page("Báo cáo được gán", body)); return
        match = re.fullmatch(r"/case/(\d+)", parsed.path)
        if match:
            case_id=int(match.group(1)); case=case_from_id(case_id); context=user_context(self.headers); legacy=session_info(self.headers)
            admin = admin_context(self.headers)
            current_context = admin or (context if case and context and report_permission(case, context, "view") else None)
            if current_context or (legacy and legacy[0]==case_id): self.send_html(case_dashboard(case_id, query, current_context)); return
            self.send_html(access_page(case_id)); return
        match = re.fullmatch(r"/case/(\d+)/finding/(\d+)", parsed.path)
        if match:
            case_id,finding_id=int(match.group(1)),int(match.group(2)); case=case_from_id(case_id); context=user_context(self.headers); legacy=session_info(self.headers)
            admin = admin_context(self.headers)
            account_context = admin or (context if case and context and report_permission(case, context, "view") else None)
            if account_context: role="reviewer"
            elif legacy and legacy[0]==case_id: role=legacy[1]
            else: self.send_html(access_page(case_id)); return
            self.send_html(_render_finding_detail(case_id,finding_id,query,role,account_context)); return
        self.send_html(html_page("404", "<section class=\"band\">Không tìm thấy trang.</section>"), HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:
        parsed = urllib.parse.urlparse(self.path)
        try:
            body=self.read_body(); parts=parse_post_parts(self.headers,body)
            if parsed.path == "/api/upload":
                if not verify_upload_key(self.headers.get("X-Upload-Key", "")): self.send_json({"ok":False,"error":"invalid upload key"},HTTPStatus.FORBIDDEN); return
                bundle=file_part(parts,"review_bundle")
                if not bundle: self.send_json({"ok":False,"error":"missing .enc"},HTTPStatus.BAD_REQUEST); return
                case_id,_,imported,source_count=import_enc_case(form_value(parts,"title"),bundle,upload_origin="remote_upload"); self.send_json({"ok":True,"case_id":case_id,"imported":imported,"source_count":source_count}); return
            if parsed.path == "/login":
                username=normalize_username(form_value(parts,"username")); password=form_value(parts,"password"); user=query_one("SELECT * FROM users WHERE username=? AND active=1",(username,))
                if not user or not verify_password(password,user["password_hash"]): self.send_html(login_page("Tên đăng nhập hoặc mật khẩu không hợp lệ."),HTTPStatus.FORBIDDEN); return
                if user["role"] == "admin":
                    token, _ = create_admin_session(int(user["id"]))
                    self.redirect("/admin", cookie=f"{ADMIN_COOKIE}={token}; Path=/; Max-Age={SESSION_TTL_SECONDS}; HttpOnly; SameSite=Strict")
                else:
                    token, _ = create_user_session(int(user["id"]))
                    self.redirect("/user", cookie=f"{USER_COOKIE}={token}; Path=/; Max-Age={SESSION_TTL_SECONDS}; HttpOnly; SameSite=Strict")
                return
            if parsed.path == "/admin/login":
                if form_value(parts,"admin_key") != get_admin_key(): self.send_html(admin_login_page("Admin key không hợp lệ."),HTTPStatus.FORBIDDEN); return
                admin_user=query_one("SELECT id FROM users WHERE username='antt' AND role='admin' LIMIT 1"); token,_=create_admin_session(admin_user["id"] if admin_user else None); self.redirect("/admin",cookie=f"{ADMIN_COOKIE}={token}; Path=/; Max-Age={SESSION_TTL_SECONDS}; HttpOnly; SameSite=Strict"); return
            if parsed.path == "/access":
                key=form_value(parts,"access_key"); requested=safe_int(form_value(parts,"case_id"),None); case=case_access_from_key(key)
                if not case or (requested and requested != int(case["id"])): self.send_html(access_page(requested,"Mã truy cập không hợp lệ."),HTTPStatus.FORBIDDEN); return
                token=create_session(int(case["id"]),str(case["access_role"])); self.redirect(f"/case/{case['id']}",cookie=f"{SESSION_COOKIE}={token}; Path=/; Max-Age={SESSION_TTL_SECONDS}; HttpOnly; SameSite=Strict"); return
            admin=admin_context(self.headers)
            if parsed.path == "/admin/users/create":
                if not admin or not csrf_valid(admin,form_value(parts,"csrf_token")): self.send_html(admin_login_page("Phiên admin hoặc CSRF không hợp lệ."),HTTPStatus.FORBIDDEN); return
                create_customer_user(form_value(parts,"username"),form_value(parts,"password"),admin.get("user_id") or 0); self.redirect("/admin"); return
            match=re.fullmatch(r"/admin/users/(\d+)/active",parsed.path)
            if match:
                if not admin or not csrf_valid(admin,form_value(parts,"csrf_token")): self.send_html(admin_login_page("Phiên admin hoặc CSRF không hợp lệ."),HTTPStatus.FORBIDDEN); return
                set_user_active(int(match.group(1)),form_value(parts,"active")=="1",admin.get("user_id") or 0); self.redirect("/admin"); return
            match=re.fullmatch(r"/admin/cases/(\d+)/(assign|revoke)",parsed.path)
            if match:
                if not admin or not csrf_valid(admin,form_value(parts,"csrf_token")): self.send_html(admin_login_page("Phiên admin hoặc CSRF không hợp lệ."),HTTPStatus.FORBIDDEN); return
                case_id=int(match.group(1)); user_id=int(form_value(parts,"user_id")); assign_report(case_id,user_id,admin.get("user_id") or 0) if match.group(2)=="assign" else revoke_report(case_id,user_id,admin.get("user_id") or 0); self.redirect("/admin"); return
            if parsed.path == "/admin/import":
                if not admin or not csrf_valid(admin,form_value(parts,"csrf_token")): self.send_html(admin_login_page("Phiên admin hoặc CSRF không hợp lệ."),HTTPStatus.FORBIDDEN); return
                bundle=file_part(parts,"review_bundle")
                if not bundle: raise RuntimeError("Thiếu file .enc.")
                import_enc_case(form_value(parts,"title"),bundle); self.redirect("/admin"); return
            match=re.fullmatch(r"/case/(\d+)/findings/bulk",parsed.path)
            if match:
                case_id=int(match.group(1)); context=user_context(self.headers); admin_context_value=admin_context(self.headers); actor=admin_context_value or context
                if not actor or not csrf_valid(actor,form_value(parts,"csrf_token")): self.send_html(access_page(case_id,"Phiên hoặc CSRF không hợp lệ."),HTTPStatus.FORBIDDEN); return
                ids=[int(part.data.decode("utf-8",errors="replace").strip()) for part in parts.get("finding_ids",[]) if part.data.decode("utf-8",errors="replace").strip().isdigit()]
                changed=bulk_update_findings(case_id,ids,form_value(parts,"status"),form_value(parts,"note"),actor)
                self.redirect(f"/case/{case_id}?updated={changed}"); return
            match=re.fullmatch(r"/case/(\d+)/finding/(\d+)/update",parsed.path)
            if match:
                case_id,finding_id=int(match.group(1)),int(match.group(2)); case=case_from_id(case_id); context=user_context(self.headers); admin_context_value=admin_context(self.headers); allowed=bool(admin_context_value and csrf_valid(admin_context_value,form_value(parts,"csrf_token"))) or bool(context and report_permission(case,context,"process") and csrf_valid(context,form_value(parts,"csrf_token")))
                if not allowed:
                    legacy=session_info(self.headers); allowed=bool(legacy and legacy[0]==case_id and legacy[1]=="reviewer")
                if not case or not allowed: self.send_html(access_page(case_id,"Phiên hoặc quyền truy cập không hợp lệ."),HTTPStatus.FORBIDDEN); return
                status=form_value(parts,"status","pending"); status=status if status in VALID_STATUSES else "pending"; note=form_value(parts,"reviewer_note","")
                if status=="benign" and not note.strip(): self.send_html(html_page("Thiếu ghi chú","<section class=\"band warn\">Bỏ qua cần có ghi chú.</section>"),HTTPStatus.BAD_REQUEST); return
                with connect_db() as conn: conn.execute("UPDATE findings SET status=?,reviewer_note=? WHERE id=? AND case_id=?",(status,note,finding_id,case_id)); conn.execute("INSERT INTO audit_log(case_id,finding_id,action,detail,created_at) VALUES(?,?,?,?,?)",(case_id,finding_id,"review_update","note_present=%s"%bool(note.strip()),now_ts()))
                self.redirect(f"/case/{case_id}/finding/{finding_id}"); return
            self.send_html(html_page("404","<section class=\"band\">Không tìm thấy endpoint.</section>"),HTTPStatus.NOT_FOUND)
        except Exception as exc:
            if parsed.path=="/api/upload": self.send_json({"ok":False,"error":"upload failed"},HTTPStatus.BAD_REQUEST)
            else: self.send_html(html_page("Lỗi",f"<section class=\"band\"><h1>Lỗi xử lý</h1><div class=\"warn\">{html.escape(str(exc))}</div></section>"),HTTPStatus.BAD_REQUEST)

    def log_message(self, fmt, *args) -> None:
        sys.stderr.write("%s - %s\n" % (self.log_date_time_string(), fmt % args))


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except Exception:
                pass

    parser = argparse.ArgumentParser(description="Web review portal for webshell scan findings")
    parser.add_argument("--host", default="127.0.0.1", help="Host bind, default 127.0.0.1")
    parser.add_argument("--port", type=int, default=8088, help="Port, default 8088")
    args = parser.parse_args(argv)

    ensure_dirs()
    init_db()
    compile_yara_rules()
    admin_url = f"http://{args.host}:{args.port}/admin"
    print("Webshell Scan Review")
    print(f"Admin URL: {admin_url}")
    print(f"YARA: {_YARA_STATUS or 'not initialized'}")
    server = ThreadingHTTPServer((args.host, args.port), ReviewHandler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping...")
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
